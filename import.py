import os
import openpyxl
import psycopg2

from config import DB_CONFIG, DATA_DIR


# ============================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================
def _load_sheet(filename):
    """Загружает активный лист Excel-файла. Возвращает (sheet, path) или (None, path)."""
    path = os.path.join(DATA_DIR, filename)
    if not os.path.exists(path):
        print(f"⚠️ Файл не найден: {path}")
        return None, path

    workbook = openpyxl.load_workbook(path)
    return workbook.active, path


# ============================================================
# ИМПОРТ РАЗМЕРОВ
# ============================================================
def load_sizes(cur):
    sheet, _ = _load_sheet("Sizes_import.xlsx")
    if sheet is None:
        return

    counter = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        size_val = float(row[0])
        cur.execute(
            "INSERT INTO sizes (size_value) VALUES (%s) ON CONFLICT (size_value) DO NOTHING;",
            (size_val,)
        )
        counter += 1

    print(f"✅ Импортировано размеров: {counter}")


# ============================================================
# ИМПОРТ ПОЛЬЗОВАТЕЛЕЙ
# ============================================================
def load_users(cur):
    sheet, _ = _load_sheet("Users_import.xlsx")
    if sheet is None:
        return

    counter = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[3] is None:
            continue

        surname, name, patronymic, login, role = row[0], row[1], row[2], row[3], row[4]
        cur.execute(
            """
            INSERT INTO users (last_name, first_name, middle_name, login, role)
            VALUES (%s, %s, %s, %s, %s)
            ON CONFLICT (login) DO NOTHING;
            """,
            (surname, name, patronymic, login, role)
        )
        counter += 1

    print(f"✅ Импортировано пользователей: {counter}")


# ============================================================
# ИМПОРТ ТОВАРОВ
# ============================================================
def load_products(cur):
    sheet, _ = _load_sheet("Products_import.xlsx")
    if sheet is None:
        return

    counter = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[3] is None:
            continue

        category, subcategory, image, name, manufacturer, description, composition, price = row
        cur.execute(
            """
            INSERT INTO products (category, subcategory, image, product_name, manufacturer, description, composition, price)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s);
            """,
            (category, subcategory, image, name, manufacturer, description, composition, price)
        )
        counter += 1

    print(f"✅ Импортировано товаров: {counter}")


# ============================================================
# ИМПОРТ СКЛАДА
# ============================================================
def load_stock(cur):
    sheet, _ = _load_sheet("Stock_Items_import.xlsx")
    if sheet is None:
        return

    counter = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        name, manufacturer, size, quantity = row
        cur.execute(
            """
            INSERT INTO stock_items (product_name, manufacturer, size, stock_quantity)
            VALUES (%s, %s, %s, %s);
            """,
            (name, manufacturer, size, quantity)
        )
        counter += 1

    print(f"✅ Импортировано складских позиций: {counter}")


# ============================================================
# ИМПОРТ ЗАКАЗОВ
# ============================================================
def load_orders(cur):
    sheet, _ = _load_sheet("Orders_import.xlsx")
    if sheet is None:
        return

    counter = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        (order_number, order_date, customer_name, category,
         product_name, manufacturer, size, quantity, unit_price) = row

        cur.execute(
            """
            INSERT INTO orders (order_number, order_date, customer_name, category, product_name,
                                manufacturer, size, quantity, unit_price)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s);
            """,
            (order_number, order_date, customer_name, category,
             product_name, manufacturer, size, quantity, unit_price)
        )
        counter += 1

    print(f"✅ Импортировано заказов: {counter}")


# ============================================================
# ГЛАВНАЯ ФУНКЦИЯ
# ============================================================
def main():
    print("=" * 50)
    print("Импорт данных в БД Чудо Обувь")
    print("=" * 50)

    connection = psycopg2.connect(**DB_CONFIG)
    connection.autocommit = False
    cur = connection.cursor()

    steps = [
        ("[1/5] Импорт размеров...", load_sizes),
        ("[2/5] Импорт пользователей...", load_users),
        ("[3/5] Импорт товаров...", load_products),
        ("[4/5] Импорт складских остатков...", load_stock),
        ("[5/5] Импорт заказов...", load_orders),
    ]

    try:
        for title, func in steps:
            print(f"\n{title}")
            func(cur)

        connection.commit()
        print("\n🎉 Все данные успешно импортированы в базу данных!")
    except Exception as err:
        connection.rollback()
        print(f"\n❌ Произошла ошибка при импорте: {err}")
    finally:
        cur.close()
        connection.close()


if __name__ == "__main__":
    main()